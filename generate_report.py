import os
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE_DIR = "/home/crei03/UTP/lab2"
SCREENSHOTS_DIR = os.path.join(BASE_DIR, "Selenium photos")


def generate_report():
    doc = Document()

    title = doc.add_heading("Reporte de Pruebas - Laboratorio 2", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Diseño y Ejecución de Pruebas Funcionales con Selenium\n").bold = True
    subtitle.add_run(f"Fecha: {datetime.now().strftime('%d/%m/%Y %H:%M')}\n")
    subtitle.add_run("Framework: pytest + Selenium WebDriver + Firefox\n")
    subtitle.add_run("Sitio: https://www.saucedemo.com/")

    doc.add_heading("Resumen de Resultados", level=1)
    doc.add_paragraph("Las 3 pruebas fueron ejecutadas exitosamente con pytest.")

    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Caso"
    hdr[1].text = "Resultado"
    hdr[2].text = "Descripción"

    rows_data = [
        ("Login exitoso", "PASS", "standard_user/secret_sauce → redirige a inventory.html"),
        ("Agregar al carrito", "PASS", "Add to cart → badge del carrito muestra 1"),
        ("Login fallido (bloqueado)", "PASS", "locked_out_user → mensaje 'locked out'"),
    ]
    for caso, res, desc in rows_data:
        row = table.add_row().cells
        row[0].text = caso
        row[1].text = res
        row[2].text = desc

    doc.add_heading("Estructura del Código", level=1)

    doc.add_heading("Arquitectura Page Object Model (POM)", level=2)

    doc.add_heading("pages/login_page.py", level=3)
    doc.add_paragraph(
        "Clase que encapsula la página de inicio de sesión:\n"
        "- open(): Navega a saucedemo.com y maximiza la ventana\n"
        "- enter_username(username): Escribe el usuario en el campo user-name\n"
        "- enter_password(password): Escribe la contraseña en el campo password\n"
        "- click_login(): Hace clic en el botón login-button\n"
        "- get_error_message(): Obtiene el texto del mensaje de error (data-test='error')\n"
        "- get_current_url(): Devuelve la URL actual para validaciones\n"
        "Todas las interacciones usan WebDriverWait con timeout de 10s."
    )

    doc.add_heading("pages/inventory_page.py", level=3)
    doc.add_paragraph(
        "Clase que encapsula la página de inventario:\n"
        "- wait_for_load(): Espera a que el contenedor inventory_container esté presente\n"
        "- is_on_inventory_page(): Valida que la URL contenga 'inventory.html'\n"
        "- add_backpack_to_cart(): Agrega Sauce Labs Backpack al carrito\n"
        "- get_cart_count(): Lee el badge shopping_cart_badge y retorna el número"
    )

    doc.add_heading("conftest.py", level=3)
    doc.add_paragraph(
        "Archivo de configuración de pytest que define:\n"
        "- Fixture driver(): Inicializa Firefox con geckodriver antes de cada test y lo cierra al finalizar\n"
        "- pytest_runtest_makereport(): Hook que captura automáticamente un screenshot cuando un test falla y lo incrusta en el reporte HTML usando pytest_html.extras"
    )

    doc.add_heading("test_tienda.py", level=3)
    doc.add_paragraph(
        "Archivo con 3 casos de prueba usando pytest:\n"
        "- test_login_exitoso(): Login con credenciales válidas, valida redirección a inventory\n"
        "- test_agregar_carrito(): Login + agregar producto al carrito + validar badge = 1\n"
        "- test_login_fallido_bloqueado(): Login con locked_out_user, valida mensaje de error 'Sorry, this user has been locked out'\n\n"
        "Cada test toma screenshots en cada paso usando la función snap() y las guarda en Selenium photos/.\n"
        "Los tests son independientes entre sí: cada uno maneja su propia sesión de login."
    )

    doc.add_heading("Reporte HTML (pytest-html)", level=3)
    doc.add_paragraph(
        "Se generó un reporte HTML interactivo usando pytest-html:\n"
        "- Comando: pytest test_tienda.py --html=reports/reporte_pruebas.html --self-contained-html\n"
        "- El reporte incluye metadatos (Python, plataforma, plugins), resultados por test, tiempos de ejecución\n"
        "- Si un test falla, el hook en conftest.py captura un screenshot automático que se incrusta en el reporte\n"
        "- Archivo generado: reports/reporte_pruebas.html (se abre en cualquier navegador)"
    )

    doc.add_heading("Capturas de Pantalla", level=1)

    screenshots = sorted(f for f in os.listdir(SCREENSHOTS_DIR) if f.endswith(".png"))

    descriptions = {
        "01_login_page": "Página de inicio de sesión de saucedemo.com",
        "02_username_entered": "Usuario 'standard_user' ingresado en el campo de usuario",
        "03_password_entered": "Contraseña ingresada en el campo de contraseña",
        "04_after_login_click": "Pantalla después de hacer clic en Login",
        "05_inventory_page": "Página de inventario después del login exitoso",
        "06_inventory_loaded": "Página de inventario cargada para prueba de carrito",
        "07_after_add_to_cart": "Botón cambiado a 'Remove' después de agregar al carrito",
        "08_cart_badge": "Badge del carrito mostrando 1 producto",
        "09_login_page_bloqueado": "Página de login para prueba de usuario bloqueado",
        "10_username_locked_entered": "Usuario 'locked_out_user' ingresado",
        "11_password_locked_entered": "Contraseña ingresada para locked_out_user",
        "12_after_locked_login": "Pantalla después de login fallido con locked_out_user",
        "13_error_message_locked": "Mensaje de error 'locked out' mostrado correctamente",
    }

    for s in screenshots:
        img_path = os.path.join(SCREENSHOTS_DIR, s)
        desc = "Captura de pantalla"
        for key, text in descriptions.items():
            if key in s:
                desc = text
                break
        doc.add_picture(img_path, width=Inches(5.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = caption.add_run(f"{s} - {desc}")
        run.font.size = Pt(9)
        run.font.italic = True

    doc.add_heading("Pipeline CI/CD (GitHub Actions)", level=1)
    doc.add_paragraph(
        "Como parte del laboratorio, se configuró un pipeline de integración continua "
        "usando GitHub Actions para ejecutar las pruebas automáticamente en la nube.\n\n"
        "Archivos creados:\n"
        "- requirements.txt: Declara las dependencias (selenium, pytest, pytest-html)\n"
        "- .github/workflows/python-tests.yml: Define el pipeline\n\n"
        "Flujo del pipeline:\n"
        "1. Trigger: Se activa en cada push o Pull Request a la rama main\n"
        "2. Entorno: ubuntu-latest (servidor Linux sin interfaz gráfica)\n"
         "3. Instalación: Firefox vía apt-get, geckodriver descargado desde GitHub Releases, dependencias Python con pip\n"
        "4. Ejecución: Las pruebas corren en modo headless (HEADLESS=true)\n"
        "5. Artefacto: El reporte HTML se guarda como artifact descargable\n\n"
        "Para usarlo, solo se debe subir la carpeta lab2/ a un repositorio de GitHub. "
        "El Action se activará automáticamente en cada push a main y el reporte "
        "HTML estará disponible en la sección 'Artifacts' de la ejecución."
    )

    doc.add_heading("Conclusiones", level=1)
    doc.add_paragraph(
        "Se implementaron y ejecutaron exitosamente 3 pruebas funcionales automatizadas "
        "en saucedemo.com utilizando:\n\n"
        "- Selenium WebDriver con Firefox como motor de automatización\n"
        "- Page Object Model (POM) para código organizado y mantenible\n"
        "- WebDriverWait (esperas explícitas) para robustez contra latencia\n"
        "- pytest + pytest-html para reportes profesionales\n"
        "- Auto-screenshots en fallos vía conftest.py\n"
        "- Capturas paso a paso para evidencia visual completa\n"
        "- Pipeline CI/CD con GitHub Actions para ejecución automática en la nube\n\n"
        "Resultados: 3/3 pruebas pasaron exitosamente."
    )

    report_path = os.path.join(BASE_DIR, "Reporte_Laboratorio2.docx")
    doc.save(report_path)
    print(f"Reporte DOCX generado: {report_path}")
    return report_path


if __name__ == "__main__":
    generate_report()
